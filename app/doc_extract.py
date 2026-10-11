"""Extract product information from unstructured files (image / PDF / Excel / CSV) with Gemini, match to SKUs,
and merge into the product master."""
from __future__ import annotations

import io
import json
import re
import time
import unicodedata

from .config import settings
from . import storage

SUPPORTED = {"jpg", "jpeg", "png", "webp", "pdf", "xlsx", "csv"}
MAX_FILES = 20
MAX_BYTES = 30 * 1024 * 1024
MAX_ROWS = 2000

EXTRACT_SCHEMA = {
    "type": "OBJECT",
    "properties": {
        "document_type": {"type": "STRING", "enum": ["spec_sheet", "catalog", "price_list", "label_photo", "manual", "other"]},
        "products": {"type": "ARRAY", "items": {"type": "OBJECT", "properties": {
            "model_number": {"type": "STRING"}, "jan": {"type": "STRING"}, "sku_hint": {"type": "STRING"},
            "title": {"type": "STRING"}, "brand": {"type": "STRING"}, "category": {"type": "STRING"},
            "price": {"type": "NUMBER", "nullable": True}, "description": {"type": "STRING"},
            "attributes": {"type": "ARRAY", "items": {"type": "OBJECT", "properties": {
                "key": {"type": "STRING"}, "label": {"type": "STRING"}, "value": {"type": "STRING"},
                "evidence": {"type": "STRING"}, "page": {"type": "INTEGER", "nullable": True}},
                "required": ["key", "label", "value", "evidence"]}},
            "confidence": {"type": "NUMBER"}},
            "required": ["model_number", "title", "attributes", "confidence"]}},
    },
    "required": ["document_type", "products"],
}


def ext_of(name: str) -> str:
    return name.lower().rsplit(".", 1)[-1] if "." in name else ""


def is_supported(name: str) -> bool:
    return ext_of(name) in SUPPORTED


def xlsx_to_text(data: bytes, max_rows: int = MAX_ROWS) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"### シート: {ws.title}")
        n = 0
        for row in ws.iter_rows(values_only=True):
            cells = ["" if c is None else str(c).replace("\t", " ").replace("\n", " ") for c in row]
            if not any(cells):
                continue
            out.append("\t".join(cells).rstrip("\t"))
            n += 1
            if n >= max_rows:
                out.append("…（以降省略）")
                break
    return "\n".join(out)


def known_keys_text(products: list[dict]) -> str:
    from .schema import ATTRS
    keys: dict[str, str] = {}
    for a in ATTRS:
        for rk in a.raw_keys:
            keys.setdefault(rk, a.label)
    for p in products:
        for k in (p.get("raw_attributes") or {}):
            keys.setdefault(k, "")
    return ", ".join(f"{k}({v})" if v else k for k, v in sorted(keys.items()))


def extract_prompt(file_name: str, keys_text: str) -> str:
    return f"""あなたは総合EC・家電リテーラーの商品マスタ整備担当です。添付の資料「{file_name}」（仕様書・カタログ・価格表・銘板写真など）から、記載されている商品ごとに商品情報を抽出してください。

## 厳守ルール
1. 資料に書かれていない値は絶対に作らない（推測・一般常識での補完は禁止）。
2. value は資料の表記を単位込みでそのまま転記する（換算・丸めをしない）。
3. evidence は根拠となる資料上の文字列を一字一句そのままコピーする（40字以内）。
4. key は英小文字のスネークケース。既存カタログのキーに該当する場合は必ずそのキーを使う。該当がなければ内容を表す英語のスネークケースを新しく作る。label は日本語の項目名。
   既存キー: {keys_text}
5. model_number は型番（品番）。jan は JAN/EAN コード（13桁）があれば。price は税込価格が明記されていれば数値で、なければ null。
6. description は資料内の商品説明文（要約しない、なければ空文字）。
7. 1つの資料に複数の商品があれば products に複数入れる（表の各行が商品なら行ごと）。商品情報がなければ products は空配列。
8. page は PDF のページ番号（1始まり）、それ以外は null。confidence は抽出全体の確からしさ 0.0〜1.0。
"""


async def extract_file(f: dict, products: list[dict]) -> dict:
    """f = {uri, name, mime}. Returns {document_type, products, input_tokens, output_tokens, latency_ms}."""
    from google.genai import types
    t0 = time.time()
    name = f.get("name") or f["uri"].rsplit("/", 1)[-1]
    ext = ext_of(name)
    mime = f.get("mime") or storage.guess_mime(name)
    prompt = extract_prompt(name, known_keys_text(products))
    if ext in ("xlsx", "csv"):
        data, _ = storage.download(f["uri"])
        if len(data) > MAX_BYTES:
            raise ValueError("ファイルが 30MB を超えています")
        if ext == "xlsx":
            text = xlsx_to_text(data)
        else:
            from .csv_ingest import decode_csv_bytes
            text = "\n".join(decode_csv_bytes(data).splitlines()[:MAX_ROWS])
        contents = [prompt + "\n## 資料の内容（表形式・タブ区切り）\n" + text[:400_000]]
    else:
        contents = [types.Part.from_uri(file_uri=f["uri"], mime_type=mime), prompt]
    if settings.offline:
        return {"document_type": "other", "products": [], "input_tokens": 0, "output_tokens": 0, "latency_ms": 1}
    from .gemini_client import client
    resp = await client().aio.models.generate_content(
        model=settings.gemini_model, contents=contents,
        config=types.GenerateContentConfig(temperature=0.0, response_mime_type="application/json",
                                           response_schema=EXTRACT_SCHEMA,
                                           thinking_config=types.ThinkingConfig(thinking_level=settings.thinking_level or "LOW")))
    from .pipeline import parse_json_text
    data = parse_json_text(resp.text)
    usage = resp.usage_metadata
    return {
        "document_type": data.get("document_type") or "other",
        "products": data.get("products") or [],
        "input_tokens": getattr(usage, "prompt_token_count", 0) or 0,
        "output_tokens": (getattr(usage, "candidates_token_count", 0) or 0) + (getattr(usage, "thoughts_token_count", 0) or 0),
        "latency_ms": int((time.time() - t0) * 1000),
    }


# ── SKU matching ─────────────────────────────────────────────────────────────

def norm_model(s: str) -> str:
    s = unicodedata.normalize("NFKC", s or "").upper()
    return re.sub(r"[^A-Z0-9]", "", s)


def _bigrams(s: str) -> set[str]:
    s = re.sub(r"\s+", "", unicodedata.normalize("NFKC", s or "").lower())
    return {s[i:i + 2] for i in range(len(s) - 1)}


def title_sim(a: str, b: str) -> float:
    A, B = _bigrams(a), _bigrams(b)
    return len(A & B) / max(1, len(A | B))


def match_sku(ex: dict, products: list[dict], limit: int = 5) -> list[dict]:
    mn = norm_model(ex.get("model_number") or "")
    jan = re.sub(r"\D", "", ex.get("jan") or "")
    hint = (ex.get("sku_hint") or "").strip()
    out = []
    for p in products:
        raw = p.get("raw_attributes") or {}
        score, why = 0.0, ""
        if hint and hint == p["sku"]:
            score, why = 1.0, "SKU一致"
        elif jan and jan in (re.sub(r"\D", "", str(raw.get("jan_code") or "")), p["sku"]):
            score, why = 0.99, "JAN一致"
        else:
            pm = norm_model(str(raw.get("model_number") or ""))
            if mn and pm and mn == pm:
                score, why = 0.98, "型番一致"
            elif mn and len(mn) >= 5 and mn in norm_model(p.get("title") or ""):
                score, why = 0.9, "タイトルに型番"
            elif mn and pm and len(mn) >= 5 and (mn.startswith(pm) or pm.startswith(mn)):
                score, why = 0.8, "型番の前方一致"
            else:
                s = title_sim(ex.get("title") or "", p.get("title") or "")
                if s >= 0.25:
                    score, why = round(0.4 + s * 0.4, 3), f"タイトル類似 {s:.0%}"
        if score:
            out.append({"sku": p["sku"], "title": p.get("title") or "", "image_url": p.get("image_url") or "",
                        "score": score, "why": why})
    out.sort(key=lambda m: -m["score"])
    return out[:limit]


# ── Merge ────────────────────────────────────────────────────────────────────

IDENTITY_FIELDS = {"title", "brand", "category"}


def _same(a, b) -> bool:
    try:
        return float(a) == float(b)
    except (TypeError, ValueError):
        return str(a).strip() == str(b).strip()


def _blank(v) -> bool:
    return v is None or str(v).strip() == ""


def merge(product: dict | None, ex: dict, mode: str, sku: str, file_name: str, file_uri: str,
          group_fn=None) -> tuple[dict, list[dict]]:
    """Return (merged_product, diff). mode: fill (only empty fields) | overwrite."""
    is_new = product is None
    base = json.loads(json.dumps(product)) if product else {
        "sku": sku, "title": "", "brand": "", "category": "", "description": "", "product_url": "", "image_url": "",
        "price": None, "group": None, "raw_attributes": {}, "doc_sources": [],
    }
    diff: list[dict] = []

    def put(field: str, label: str, before, after, setter):
        if _blank(after):
            return
        if is_new:
            action = "new"
        elif _blank(before):
            action = "add"
        elif _same(before, after):
            action = "keep"
        elif field in IDENTITY_FIELDS:
            action = "keep"   # never replace an existing product name/brand/category with document wording
        else:
            action = "overwrite" if mode == "overwrite" else "keep"
        if action in ("new", "add", "overwrite"):
            setter(after)
        diff.append({"field": field, "label": label, "before": "" if before is None else str(before),
                     "after": str(after), "action": action})

    put("title", "商品名", base.get("title"), ex.get("title"), lambda v: base.__setitem__("title", v))
    put("brand", "ブランド", base.get("brand"), ex.get("brand"), lambda v: base.__setitem__("brand", v))
    put("category", "カテゴリ", base.get("category"), ex.get("category"), lambda v: base.__setitem__("category", v))
    put("price", "価格", base.get("price"), ex.get("price"), lambda v: base.__setitem__("price", float(v)))
    put("description", "説明文", base.get("description"), ex.get("description"),
        lambda v: base.__setitem__("description", v))
    raw = base.setdefault("raw_attributes", {})
    if ex.get("model_number"):
        put("attr:model_number", "型番", raw.get("model_number"), ex["model_number"],
            lambda v: raw.__setitem__("model_number", v))
    docs = [d for d in base.get("doc_sources") or [] if d.get("file_uri") != file_uri]
    for a in ex.get("attributes") or []:
        k = re.sub(r"[^a-z0-9_]", "_", str(a.get("key") or "").lower()).strip("_")
        if not k:
            continue
        before = raw.get(k)
        n_before = len(diff)
        put(f"attr:{k}", a.get("label") or k, before, a.get("value"), lambda v, k=k: raw.__setitem__(k, v))
        if len(diff) > n_before and diff[-1]["action"] in ("new", "add", "overwrite"):
            docs.append({"file_name": file_name, "file_uri": file_uri, "page": a.get("page"), "key": k,
                         "label": a.get("label") or k, "value": a.get("value"), "evidence": a.get("evidence") or ""})
    base["doc_sources"] = docs
    if not base.get("title"):
        base["title"] = ex.get("model_number") or sku
    if not base.get("group"):
        from .schema import category_group
        base["group"] = category_group((base.get("category") or "") + " " + (base.get("title") or ""))
    base["source"] = "doc"
    return base, diff
